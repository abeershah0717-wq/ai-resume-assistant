"""
ATS Resume Checker
------------------
Upload a resume (PDF or DOCX), optionally paste a job description, and get:
  * an estimated ATS score (0-100) with a category breakdown
  * strengths, weaknesses and missing keywords
  * concrete, prioritised improvement suggestions with example rewrites

UI: Streamlit   |   AI: Google Gemini Flash (via the `google-genai` SDK)
"""

from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import streamlit as st

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

# "gemini-flash-latest" is an alias that always points at Google's newest Flash
# model, so the app keeps working when older versions are retired.
# You can override it in the sidebar or with the GEMINI_MODEL env var / secret.
DEFAULT_MODEL = "gemini-flash-latest"

MAX_FILE_MB = 5
MAX_RESUME_CHARS = 30_000  # protects against huge files / token blow-ups
MIN_RESUME_CHARS = 150     # below this the file is probably scanned / empty

# Category name -> maximum points. These add up to 100.
CATEGORIES: dict[str, int] = {
    "Keywords & Skills Match": 30,
    "Formatting & ATS Parseability": 20,
    "Work Experience & Impact": 25,
    "Structure & Completeness": 15,
    "Language & Readability": 10,
}

SYSTEM_INSTRUCTION = (
    "You are an expert technical recruiter and Applicant Tracking System (ATS) "
    "specialist. You evaluate resumes strictly and honestly. You never invent "
    "facts that are not in the resume. You always answer with valid JSON only."
)


# --------------------------------------------------------------------------- #
# File reading
# --------------------------------------------------------------------------- #

def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - depends on file
            raise ValueError("This PDF is password protected.") from exc
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep content inside tables (two-column layouts etc.)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    """Return cleaned plain text from a PDF or DOCX file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        raw = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        raw = extract_text_from_docx(data)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")

    text = re.sub(r"[ \t]+", " ", raw)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# --------------------------------------------------------------------------- #
# Prompt + response handling
# --------------------------------------------------------------------------- #

def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        jd_block = (
            "A TARGET JOB DESCRIPTION is provided. Judge keyword match, skills "
            "and relevance against it.\n\n=== JOB DESCRIPTION ===\n" + jd[:8000]
        )
    else:
        jd_block = (
            "No job description was provided. Judge the resume against general "
            "ATS best practices and the role the resume appears to target."
        )

    categories_text = "\n".join(f'- "{k}" (max {v})' for k, v in CATEGORIES.items())

    return f"""Analyse the resume below as an ATS would, then give an honest score.

{jd_block}

=== RESUME ===
{resume_text[:MAX_RESUME_CHARS]}

=== SCORING RUBRIC (total 100) ===
{categories_text}

Be strict: a typical resume scores 55-75; 90+ is rare. Score each category from
0 up to its max. "overall_score" must equal the sum of the category scores.

Return ONLY a JSON object with exactly this shape (no markdown, no commentary):
{{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall verdict>",
  "detected_role": "<job title the resume appears to target>",
  "category_scores": [
    {{"category": "<one of the category names above>", "score": <int>, "comment": "<one sentence>"}}
  ],
  "strengths": ["<short bullet>", "..."],
  "weaknesses": ["<short bullet>", "..."],
  "matched_keywords": ["<keyword found in resume>", "..."],
  "missing_keywords": ["<important keyword NOT in resume>", "..."],
  "formatting_issues": ["<ATS parsing/formatting problem>", "..."],
  "improvements": [
    {{
      "priority": "High" | "Medium" | "Low",
      "section": "<resume section, e.g. Experience>",
      "issue": "<what is wrong>",
      "suggestion": "<what to do>",
      "example": "<a rewritten example line using ONLY facts from the resume, or empty string>"
    }}
  ]
}}

Rules:
- Provide 5-10 improvements, ordered High -> Low priority.
- For "example", never invent employers, numbers or tools. If a metric is
  missing, use a placeholder like [X%] so the user can fill it in.
- Keep every list item concise.
"""


def _to_int(value: Any, default: int = 0) -> int:
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return default


def _str_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def parse_model_json(raw: str) -> dict:
    """Extract a JSON object from a model reply (tolerates ```json fences)."""
    if not raw or not raw.strip():
        raise ValueError("The model returned an empty response.")
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("The model reply did not contain JSON.")
        return json.loads(text[start : end + 1])


def normalize_result(data: dict) -> dict:
    """Validate and clean the model output so the UI never crashes on it."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected response format from the model.")

    # Category scores: clamp each to its maximum, then derive the overall score
    # ourselves so the numbers always add up.
    given = {}
    for item in data.get("category_scores") or []:
        if isinstance(item, dict) and item.get("category") in CATEGORIES:
            given[item["category"]] = item

    category_scores = []
    for name, max_pts in CATEGORIES.items():
        item = given.get(name, {})
        score = max(0, min(max_pts, _to_int(item.get("score"), 0)))
        category_scores.append(
            {
                "category": name,
                "score": score,
                "max": max_pts,
                "comment": str(item.get("comment", "")).strip(),
            }
        )

    if given:
        overall = sum(c["score"] for c in category_scores)
    else:  # model skipped the breakdown - fall back to its own number
        overall = max(0, min(100, _to_int(data.get("overall_score"), 0)))

    improvements = []
    for imp in data.get("improvements") or []:
        if not isinstance(imp, dict):
            continue
        priority = str(imp.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(imp.get("section", "General")).strip() or "General",
                "issue": str(imp.get("issue", "")).strip(),
                "suggestion": str(imp.get("suggestion", "")).strip(),
                "example": str(imp.get("example", "")).strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    return {
        "overall_score": overall,
        "summary": str(data.get("summary", "")).strip(),
        "detected_role": str(data.get("detected_role", "")).strip(),
        "category_scores": category_scores,
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "matched_keywords": _str_list(data.get("matched_keywords")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "formatting_issues": _str_list(data.get("formatting_issues")),
        "improvements": improvements,
    }


def analyze_resume(
    api_key: str, model: str, resume_text: str, job_description: str = ""
) -> dict:
    """Call Gemini and return a validated analysis dict."""
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_INSTRUCTION,
        response_mime_type="application/json",
        temperature=0.2,
    )
    prompt = build_prompt(resume_text, job_description)

    last_error: Exception | None = None
    for _ in range(2):  # one retry if the JSON comes back malformed
        response = client.models.generate_content(
            model=model, contents=prompt, config=config
        )
        try:
            return normalize_result(parse_model_json(response.text))
        except (ValueError, json.JSONDecodeError) as exc:
            last_error = exc
    raise ValueError(f"Could not read the AI response ({last_error}). Please retry.")


def friendly_api_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "The Gemini API key was rejected. Please check that it is correct."
    if "429" in low or "quota" in low or "resource_exhausted" in low:
        return "Gemini rate limit or quota reached. Wait a minute and try again."
    if "404" in low or "not found" in low:
        return "That model name was not found. Try 'gemini-flash-latest' in the sidebar."
    return f"Something went wrong while calling Gemini: {msg}"


def build_report_text(result: dict) -> str:
    """Plain-text report for the download button."""
    lines = [
        "ATS RESUME REPORT",
        "=" * 40,
        f"Overall ATS score: {result['overall_score']}/100",
        f"Detected role: {result['detected_role'] or 'n/a'}",
        "",
        result["summary"],
        "",
        "CATEGORY BREAKDOWN",
    ]
    for c in result["category_scores"]:
        lines.append(f"- {c['category']}: {c['score']}/{c['max']} - {c['comment']}")
    for title, key in (
        ("STRENGTHS", "strengths"),
        ("WEAKNESSES", "weaknesses"),
        ("MISSING KEYWORDS", "missing_keywords"),
        ("FORMATTING ISSUES", "formatting_issues"),
    ):
        lines += ["", title] + [f"- {x}" for x in result[key]]
    lines += ["", "IMPROVEMENTS"]
    for i, imp in enumerate(result["improvements"], 1):
        lines.append(f"{i}. [{imp['priority']}] {imp['section']}: {imp['issue']}")
        lines.append(f"   Fix: {imp['suggestion']}")
        if imp["example"]:
            lines.append(f"   Example: {imp['example']}")
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Streamlit UI
# --------------------------------------------------------------------------- #

def get_secret(name: str) -> str:
    """Read from Streamlit secrets, then environment variables."""
    try:
        value = st.secrets.get(name, "")
    except Exception:  # no secrets.toml present
        value = ""
    return str(value or os.environ.get(name, "")).strip()


def score_label(score: int) -> tuple[str, str]:
    if score >= 80:
        return "Excellent", "🟢"
    if score >= 65:
        return "Good - room to improve", "🟡"
    if score >= 50:
        return "Needs work", "🟠"
    return "Poor - major fixes needed", "🔴"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    label, icon = score_label(score)

    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score} / 100")
        st.markdown(f"### {icon} {label}")
        if result["detected_role"]:
            st.caption(f"Detected target role: **{result['detected_role']}**")
    with col2:
        st.progress(score / 100)
        st.write(result["summary"])

    st.subheader("Score breakdown")
    for c in result["category_scores"]:
        st.write(f"**{c['category']}** - {c['score']}/{c['max']}")
        st.progress(c["score"] / c["max"])
        if c["comment"]:
            st.caption(c["comment"])

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for item in result["strengths"] or ["No specific strengths returned."]:
            st.markdown(f"- {item}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for item in result["weaknesses"] or ["No specific weaknesses returned."]:
            st.markdown(f"- {item}")

    st.subheader("🔑 Keywords")
    k1, k2 = st.columns(2)
    with k1:
        st.markdown("**Found in your resume**")
        st.write(", ".join(result["matched_keywords"]) or "-")
    with k2:
        st.markdown("**Missing (consider adding where truthful)**")
        st.write(", ".join(result["missing_keywords"]) or "-")

    if result["formatting_issues"]:
        st.subheader("📄 Formatting issues")
        for item in result["formatting_issues"]:
            st.markdown(f"- {item}")

    st.subheader("🛠️ Recommended improvements")
    badge = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for imp in result["improvements"]:
        title = f"{badge[imp['priority']]} {imp['priority']} · {imp['section']}"
        with st.expander(title, expanded=imp["priority"] == "High"):
            st.markdown(f"**Issue:** {imp['issue']}")
            st.markdown(f"**How to fix:** {imp['suggestion']}")
            if imp["example"]:
                st.markdown("**Example rewrite:**")
                st.code(imp["example"], language=None)

    st.download_button(
        "⬇️ Download report (.txt)",
        data=build_report_text(result),
        file_name="ats_resume_report.txt",
        mime="text/plain",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.write(
        "Upload your resume to get an estimated ATS score and specific tips to "
        "improve it. Add a job description for a tailored keyword check."
    )

    with st.sidebar:
        st.header("Settings")
        secret_key = get_secret("GEMINI_API_KEY")
        if secret_key:
            st.success("API key loaded from secrets.")
            api_key = secret_key
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            ).strip()
        model = st.text_input(
            "Gemini model", value=get_secret("GEMINI_MODEL") or DEFAULT_MODEL
        ).strip()
        st.caption(
            "The score is an AI-based estimate. Real ATS software varies, so use "
            "it as guidance, not a guarantee."
        )

    uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    job_description = st.text_area(
        "Job description (optional but recommended)",
        height=160,
        placeholder="Paste the job posting here to check keyword match...",
    )

    if st.button("Analyze resume", type="primary"):
        if uploaded is None:
            st.warning("Please upload a resume first.")
            return
        if not api_key:
            st.warning("Please enter your Gemini API key in the sidebar.")
            return
        if not model:
            st.warning("Please enter a Gemini model name in the sidebar.")
            return

        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, data)
        except Exception as exc:
            st.error(f"Could not read this file: {exc}")
            return

        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned "
                "image, an ATS can't read it either - export a text-based PDF or "
                "DOCX and try again."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(
                    api_key, model, text, job_description
                )
        except ValueError as exc:
            st.error(str(exc))
            return
        except Exception as exc:
            st.error(friendly_api_error(exc))
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
